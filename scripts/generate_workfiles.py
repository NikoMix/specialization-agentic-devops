"""Generate downloadable workfiles for the Agentic DevOps specialization repo.

Writes 14 files under static/templates/{engagement,deliverables,audit}/.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PInches
from pptx.dml.color import RGBColor as PRGBColor
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent / "static" / "templates"
ENG = ROOT / "engagement"
DEL = ROOT / "deliverables"
AUD = ROOT / "audit"
for d in (ENG, DEL, AUD):
    d.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(0x0B, 0x2E, 0x5C)
PNAVY = PRGBColor(0x0B, 0x2E, 0x5C)
GREY = RGBColor(0x55, 0x55, 0x55)
HEADER_FILL = PatternFill("solid", fgColor="0B2E5C")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(border_style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_LIST = '"Not started,In progress,Blocked,Complete,N/A"'


def docx_cover(doc, title, subtitle):
    section = doc.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(32)
    run.font.color.rgb = NAVY
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    run.italic = True
    run.font.size = Pt(14)
    run.font.color.rgb = GREY
    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Agentic DevOps - Advanced Specialization")
    run.font.size = Pt(11)
    run.font.color.rgb = GREY
    doc.add_page_break()
    doc.add_heading("Document control", level=1)
    tbl = doc.add_table(rows=4, cols=2)
    tbl.style = "Light Grid Accent 1"
    for i, (k, v) in enumerate([
        ("Customer", "<customer name>"),
        ("Version", "0.1 (draft)"),
        ("Authors", "<author>"),
        ("Reviewers", "<reviewer>"),
    ]):
        tbl.rows[i].cells[0].text = k
        tbl.rows[i].cells[1].text = v
    doc.add_paragraph()
    doc.add_heading("Table of contents", level=1)
    p = doc.add_paragraph()
    r = p.add_run("[Generate TOC in Word: References -> Table of Contents]")
    r.italic = True
    r.font.color.rgb = GREY
    doc.add_page_break()


def docx_section(doc, heading, prompts):
    doc.add_heading(heading, level=1)
    for prompt in prompts:
        p = doc.add_paragraph()
        r = p.add_run(prompt)
        r.italic = True
        r.font.color.rgb = GREY
    doc.add_paragraph()


def docx_table(doc, header, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        cell = tbl.rows[0].cells[i]
        cell.text = h
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            tbl.rows[ri].cells[ci].text = val
    doc.add_paragraph()


def build_qualification_questionnaire():
    doc = Document()
    docx_cover(doc, "Customer Qualification Questionnaire", "Agentic DevOps engagement")
    sections = [
        ("1. Today's source control", [
            "What is your primary SCM (GitHub, Azure DevOps, GitLab, Bitbucket)?",
            "Any secondary systems? Plans to consolidate?",
            "How many active repositories?",
        ]),
        ("2. Today's CI/CD", [
            "How many pipelines per team on average?",
            "Median PR cycle time (open -> merge)?",
            "Green-build rate (last 30 days)?",
        ]),
        ("3. Identity & access", [
            "Identity provider (Entra ID / Okta / other)?",
            "Is GitHub Enterprise Managed Users (EMU) in use?",
            "Is audit log streaming configured? Where to?",
        ]),
        ("4. Security posture", [
            "Is GitHub Advanced Security enabled? Which features?",
            "How many open secret scanning findings?",
            "Is Dependabot configured? Auto-merge policy?",
        ]),
        ("5. AI policy", [
            "Is GitHub Copilot allowed?",
            "Which models / providers are approved?",
            "Which data classifications can be sent to AI tools?",
        ]),
        ("6. Developer environments", [
            "Are Dev Box / Codespaces in use?",
            "Local-only or restricted environment policy?",
        ]),
        ("7. Migration appetite", [
            "Is Azure DevOps -> GitHub migration on the roadmap?",
            "Timeline and budget?",
            "Named owner on the customer side?",
        ]),
        ("8. Outcomes wanted", [
            "Top 3 business outcomes in priority order.",
            "How will success be measured at engagement close?",
        ]),
    ]
    for h, prompts in sections:
        docx_section(doc, h, prompts)
    doc.add_heading("Qualification scorecard", level=1)
    docx_table(doc, ["Section", "Score (0-3)", "Rationale"], [[s[0], "", ""] for s in sections])
    p = doc.add_paragraph()
    r = p.add_run("Total: __ / 24. Recommendation: Go (17-24) / Defer (9-16) / No-go (0-8).")
    r.bold = True
    doc.save(ENG / "qualification-questionnaire.docx")


def build_definition_of_done():
    doc = Document()
    docx_cover(doc, "Definition of Done", "Agentic DevOps engagement")
    docx_section(doc, "1. Purpose", ["Lock acceptance criteria at end of discovery. Re-confirm at hypercare exit."])
    docx_section(doc, "2. Per-deliverable acceptance criteria", ["Tailor rows below to engagement scope."])
    docx_table(doc, ["Deliverable", "Acceptance criteria", "Owner", "Signed"], [
        ["Discovery report + scorecard", "Customer sponsor acknowledged in writing", "<sponsor>", ""],
        ["HLD", "Reviewed by customer architect; no open major issues", "<architect>", ""],
        ["LLD", "Per-repo + per-team config; CODEOWNER mapping signed off", "<lead engineer>", ""],
        ["Pilot enablement", "Pilot team passes success metrics for 2 sprints", "<pilot lead>", ""],
        ["Scale-out", "All teams on new platform; baseline metrics held or improved", "<platform lead>", ""],
        ["Runbook", "Dry-run incident executed by customer team", "<platform lead>", ""],
        ["KT plan", "1-hour skills check passed per module", "<platform lead>", ""],
        ["Hypercare plan", "30/60/90 SLOs met or formally accepted", "<sponsor>", ""],
    ])
    docx_section(doc, "3. Engagement-level DoD", [
        "All deliverables accepted, success metrics met or formally accepted, risks transferred, DoD signed, final invoice raised, lessons-learned issue opened.",
    ])
    doc.save(ENG / "definition-of-done.docx")


def build_deliverable_template(filename, title, sections):
    doc = Document()
    docx_cover(doc, title, "Agentic DevOps engagement - customer deliverable template")
    for heading, prompts in sections:
        docx_section(doc, heading, prompts)
    doc.save(DEL / filename)


def build_hld():
    build_deliverable_template("hld-template.docx", "High-Level Design (HLD)", [
        ("1. Executive summary", ["One page: outcomes, scope, key decisions."]),
        ("2. Business context", ["Pull from the discovery report."]),
        ("3. Current state", ["Topology diagram + pain points."]),
        ("4. Target state - platform topology", ["RA-1 (control plane) plus chosen runner / agent topologies."]),
        ("5. Identity & access", ["Entra ID, SSO, EMU decision, audit log streaming."]),
        ("6. Security architecture", ["GHAS posture, supply chain, agent permission model."]),
        ("7. Operational model", ["RACI, on-call, change management."]),
        ("8. Non-functional requirements", ["Performance, reliability, cost envelope."]),
        ("9. Open questions & risks", ["Tracked through to LLD."]),
    ])


def build_lld():
    build_deliverable_template("lld-template.docx", "Low-Level Design (LLD)", [
        ("1. Repo bootstrap", ["Default branch, protection, required workflows, CODEOWNERS, labels."]),
        ("2. Environments", ["Dev / Test / Prod; OIDC federated credential per env."]),
        ("3. Azure role assignments", ["Resource group, role, principal, justification."]),
        ("4. Key Vault scopes", ["Secret naming, access policies / RBAC, rotation cadence."]),
        ("5. GHAS configuration", ["Code scanning, secret scanning push protection, Dependabot."]),
        ("6. Runner configuration", ["ARC namespace, runner image, scaling rules, allow-listed egress."]),
        ("7. MCP server manifests", ["Per-server: tools, permissions, secrets, endpoint, audit."]),
        ("8. GitHub App manifests", ["Per-app: permissions, events, install scope, secret model."]),
        ("9. Pipeline catalog", ["Reusable workflows + custom actions in use."]),
        ("10. Validation plan", ["Smoke tests, security tests, rollback plan."]),
    ])


def build_runbook():
    build_deliverable_template("runbook-template.docx", "Platform Runbook", [
        ("1. Recurring operational procedures", ["List per-task SOPs with cadence."]),
        ("2. OIDC federated credential review (quarterly)", ["Steps to enumerate, verify, rotate."]),
        ("3. Self-hosted runner image patching (monthly)", ["Pipeline trigger + verification."]),
        ("4. ARC controller upgrade (quarterly)", ["Test in non-prod first."]),
        ("5. Copilot license review (quarterly)", ["Active vs inactive seats; reclaim."]),
        ("6. GHAS alert review (weekly)", ["High severity triage SLA."]),
        ("7. Branch protection audit (monthly)", ["Drift detection."]),
        ("8. Audit log integrity check (monthly)", ["Stream completeness verification."]),
        ("9. MCP server token rotation", ["90-day default; per-server policy."]),
        ("10. Incident response", ["Sev rubric + first-responder action per incident type."]),
    ])


def build_kt_plan():
    build_deliverable_template("kt-plan-template.docx", "Knowledge Transfer (KT) Plan", [
        ("1. Customer platform team roster", ["Name, role, GitHub handle, baseline skill 1-5."]),
        ("2. Module map", ["8 modules from GitHub admin -> Audit log + SIEM."]),
        ("3. RACI per module", ["Responsible, accountable, consulted, informed."]),
        ("4. Milestones (Shadow -> Co-pilot -> Reverse-shadow -> Independent)", ["1 sprint per stage per module."]),
        ("5. Skills-check rubric per module", ["1-hour practical test; pass/fail criteria."]),
        ("6. Schedule", ["Calendar view of shadow / reverse-shadow stages."]),
        ("7. Sign-off", ["Per-module skills-check result + customer sponsor signature."]),
    ])


def build_hypercare_plan():
    build_deliverable_template("hypercare-plan-template.docx", "Hypercare Plan", [
        ("1. Purpose & scope", ["30/60/90 partner-on-standby phase."]),
        ("2. SLOs by phase", ["Response time, pipeline health, GHAS, Copilot adoption, customer self-resolve %."]),
        ("3. Hypercare review cadence", ["Weekly review + 30/60/90 scorecards."]),
        ("4. Escalation paths", ["Sev 1 / Sev 2 contact tree."]),
        ("5. Exit gate", ["Checklist: SLOs met, no open Sev 1/2, KT confidence >= 4/5, lessons-learned raised, invoice."]),
    ])


def pptx_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    for shape in slide.placeholders:
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                run.font.color.rgb = PNAVY


def pptx_bullet_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.text = bullets[0]
    for b in bullets[1:]:
        p = body.add_paragraph()
        p.text = b
        p.level = 0


def build_offering_one_pager():
    prs = Presentation()
    prs.slide_width = PInches(13.333)
    prs.slide_height = PInches(7.5)
    pptx_title_slide(prs, "Agentic DevOps", "Offering one-pager - <customer>")
    pptx_bullet_slide(prs, "Outcomes", [
        "Secure SDLC: GHAS adopted; high-severity findings -> 0",
        "AI-assisted SDLC: Copilot Business/Enterprise + coding agent in production",
        "Pipeline modernisation: Actions on ARC runners; OIDC to Azure",
        "Toolchain consolidation: Azure DevOps -> GitHub migration if applicable",
    ])
    pptx_bullet_slide(prs, "Scope", [
        "In: GHEC tenancy design, GHAS rollout, Copilot rollout, runner platform, agent governance, ADO->GH migration",
        "Out: application refactor, end-user training beyond pilot teams",
    ])
    pptx_bullet_slide(prs, "Deliverables", [
        "Discovery report + WAF scorecard",
        "HLD + LLD",
        "Pilot enablement (one team, 2 sprints)",
        "Scale-out to 3-5 teams",
        "Runbook + KT plan + Hypercare plan",
    ])
    pptx_bullet_slide(prs, "Prerequisites", [
        "Customer sponsor + platform lead identified",
        "Read-only admin access to GHEC / Azure DevOps",
        "License entitlement: Copilot + GHAS + Actions minutes",
    ])
    pptx_bullet_slide(prs, "Timeline & price band", [
        "6-12 weeks depending on team count",
        "Indicative: <price band> - confirmed after qualification",
    ])
    pptx_bullet_slide(prs, "Next steps", [
        "Run the Customer Qualification Questionnaire",
        "Book the 2-day Discovery Workshop",
    ])
    prs.save(ENG / "offering-one-pager.pptx")


def build_discovery_workshop_deck():
    prs = Presentation()
    prs.slide_width = PInches(13.333)
    prs.slide_height = PInches(7.5)
    pptx_title_slide(prs, "Agentic DevOps - Discovery Workshop", "<customer> - Day 1-2")
    pptx_bullet_slide(prs, "Agenda - Day 1", [
        "09:00 Welcome + outcomes",
        "09:30 Stakeholder map",
        "10:30 Current-state SDLC walkthrough",
        "13:30 Pipeline + GHAS baseline tour",
        "15:00 Identity & access review",
        "16:00 Copilot adoption signals",
    ])
    pptx_bullet_slide(prs, "Agenda - Day 2", [
        "09:00 Target state - reference architecture",
        "10:30 Agent permission scoping",
        "13:30 Engagement plan draft",
        "15:00 Definition of done",
        "16:00 Next steps + open actions",
    ])
    pptx_bullet_slide(prs, "Outcomes statement", [
        "Top 3 business outcomes (priority order)",
        "Measurable success metrics per outcome",
        "Sponsor signature",
    ])
    pptx_bullet_slide(prs, "Stakeholder map", [
        "Sponsor / decision maker",
        "Platform lead",
        "Security lead",
        "Engineering manager(s) for pilot teams",
        "Procurement / commercial",
    ])
    pptx_bullet_slide(prs, "Current-state SDLC", [
        "SCM topology", "CI/CD topology", "Identity + access", "Security tooling", "Pain points",
    ])
    pptx_bullet_slide(prs, "Baseline metrics scorecard", [
        "PR cycle time", "Green-build rate", "GHAS open findings by severity", "Copilot active-user rate",
    ])
    pptx_bullet_slide(prs, "Target state - RA pick", [
        "RA-1 GitHub control plane (always)",
        "RA-2 Self-hosted runners on Azure (if egress / residency)",
        "RA-3 Copilot coding agent + MCP servers (if scope)",
        "RA-4 ADO->GH migration (if scope)",
        "RA-5 Secure SDLC reference",
        "RA-6 Agentic release pipeline",
    ])
    pptx_bullet_slide(prs, "Agent permission scoping", [
        "Identity = dedicated GitHub App (not user PAT)",
        "Per-repo scope, read-default",
        "Tool allow-list per MCP server",
        "Audit trail -> Sentinel",
    ])
    pptx_bullet_slide(prs, "Engagement plan v0.1", [
        "Workstreams + timeline", "RACI", "Risks + assumptions", "Acceptance criteria",
    ])
    pptx_bullet_slide(prs, "Definition of Done", [
        "Per-deliverable acceptance criteria", "Engagement-level checklist", "Sign-off process",
    ])
    pptx_bullet_slide(prs, "Next steps", [
        "Discovery report + scorecard within 5 business days",
        "HLD kick-off date",
        "Open actions register",
    ])
    prs.save(ENG / "discovery-workshop-deck.pptx")


def style_header_row(ws, header):
    for i, h in enumerate(header, start=1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.freeze_panes = "A2"
    for i in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 28


def add_status_dropdown(ws, column_letter, last_row=200):
    dv = DataValidation(type="list", formula1=STATUS_LIST, allow_blank=True)
    dv.add(f"{column_letter}2:{column_letter}{last_row}")
    ws.add_data_validation(dv)


def build_discovery_workshop_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Stakeholders")
    style_header_row(ws, ["Name", "Role", "Org", "Email", "RACI", "Notes"])
    ws = wb.create_sheet("Current state")
    style_header_row(ws, ["Area", "Today", "Pain points", "Owner"])
    for i, area in enumerate(["SCM", "CI/CD", "Identity", "Security", "Dev environments", "AI policy"], start=2):
        ws.cell(row=i, column=1, value=area)
    ws = wb.create_sheet("Baseline metrics")
    style_header_row(ws, ["Metric", "Today", "Target", "Source"])
    metrics = [
        ("Median PR cycle time", "", "< 1 day", "GHEC API"),
        ("Green-build rate (30d)", "", ">= 95%", "Actions API"),
        ("GHAS high-severity open", "", "0", "GHAS"),
        ("Copilot active-user rate", "", ">= 70%", "Copilot admin"),
        ("Deployment frequency", "", "Daily+", "Actions API"),
    ]
    for i, m in enumerate(metrics, start=2):
        for j, v in enumerate(m, start=1):
            ws.cell(row=i, column=j, value=v)
    ws = wb.create_sheet("Action register")
    style_header_row(ws, ["#", "Action", "Owner", "Due", "Status", "Notes"])
    add_status_dropdown(ws, "E")
    wb.save(ENG / "discovery-workshop-workbook.xlsx")


def build_waf_assessment():
    wb = Workbook()
    wb.remove(wb.active)
    pillars = {
        "Operational Excellence": [
            ("OE.1", "Every change deployable via a pipeline"),
            ("OE.2", "Median PR cycle time"),
            ("OE.3", "SDLC observability in place"),
            ("OE.4", "Runbooks documented + tested"),
            ("OE.5", "Incident retro is a standing practice"),
        ],
        "Security": [
            ("SEC.1", "Branch protection enforced on main"),
            ("SEC.2", "Secrets in Key Vault via OIDC"),
            ("SEC.3", "GHAS code scanning with severity gating"),
            ("SEC.4", "Supply chain provenance signed"),
            ("SEC.5", "Agent / Copilot access scoped"),
        ],
        "Reliability": [
            ("REL.1", "Runner pools sized + auto-scaled"),
            ("REL.2", "Pipeline retries idempotent"),
            ("REL.3", "Disaster recovery for platform"),
        ],
        "Cost Optimization": [
            ("COST.1", "Runner minute cost tracked per team"),
            ("COST.2", "Copilot seats reviewed quarterly"),
            ("COST.3", "Agent token spend monitored"),
        ],
    }
    ws = wb.create_sheet("Summary")
    style_header_row(ws, ["Pillar", "Avg score", "Top improvement"])
    for i, p in enumerate(pillars.keys(), start=2):
        ws.cell(row=i, column=1, value=p)
    for pillar, qs in pillars.items():
        ws = wb.create_sheet(pillar[:31])
        style_header_row(ws, ["#", "Question", "Score (1-5)", "Evidence", "Improvement"])
        for i, (num, q) in enumerate(qs, start=2):
            ws.cell(row=i, column=1, value=num)
            ws.cell(row=i, column=2, value=q)
        dv = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
        dv.add(f"C2:C{1 + len(qs)}")
        ws.add_data_validation(dv)
    wb.save(ENG / "waf-assessment.xlsx")


def build_assessment_platform_inputs():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Inputs - pre-baseline")
    style_header_row(ws, ["Input", "Value", "Source", "Owner"])
    for i, row in enumerate([
        ("Org-level GitHub metrics export", "", "GHEC admin", ""),
        ("Pipeline metrics (90d)", "", "GHEC / ADO", ""),
        ("License inventory (Copilot, GHAS, ADO)", "", "Billing admin", ""),
        ("Network topology summary", "", "Customer architect", ""),
    ], start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)
    ws = wb.create_sheet("Inputs - live discovery")
    style_header_row(ws, ["Topic", "Finding", "Severity", "Action"])
    ws = wb.create_sheet("DevOps Maturity")
    style_header_row(ws, ["Practice", "Current level (1-5)", "Target level (1-5)", "Notes"])
    for i, p in enumerate(["Source control", "CI", "CD", "Testing", "Security", "Observability", "Collaboration"], start=2):
        ws.cell(row=i, column=1, value=p)
    ws = wb.create_sheet("Copilot Readiness")
    style_header_row(ws, ["Question", "Answer", "Gap / action"])
    for i, q in enumerate([
        "License entitlement in place?",
        "Network egress to api.githubcopilot.com confirmed?",
        "Content exclusions configured?",
        "Acceptable-use policy published?",
        "Pilot teams identified?",
    ], start=2):
        ws.cell(row=i, column=1, value=q)
    ws = wb.create_sheet("GHAS Posture")
    style_header_row(ws, ["Feature", "Enabled (Y/N)", "Coverage %", "Owner"])
    for i, f in enumerate(["Code scanning", "Secret scanning", "Push protection", "Dependabot alerts", "Dependabot updates"], start=2):
        ws.cell(row=i, column=1, value=f)
    wb.save(ENG / "assessment-platform-inputs.xlsx")


def build_evidence_tracker():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Evidence tracker")
    style_header_row(ws, ["Module", "Control #", "Title", "Owner", "Evidence link", "Reviewer", "Status", "Last updated", "Notes"])
    add_status_dropdown(ws, "G")
    rows = [
        ("A", "A.1.1", "Organizational data"),
        ("A", "A.1.2", "Financial documentation"),
        ("A", "A.2.1", "Service delivery methodology"),
        ("A", "A.2.2", "Quality management"),
        ("A", "A.3.1", "Customer satisfaction"),
        ("A", "A.3.2", "Complaint handling"),
        ("A", "A.3.3", "Security & privacy"),
        ("B", "B.1.1", "Agentic DevOps capability"),
        ("B", "B.2.1", "ACR performance"),
        ("B", "B.2.2", "Customer diversity"),
        ("B", "B.3.1", "Certifications"),
        ("B", "B.4.1", "Audit readiness"),
        ("B", "B.4.2", "Partner onboarding"),
    ]
    for i, (m, n, t) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=m)
        ws.cell(row=i, column=2, value=n)
        ws.cell(row=i, column=3, value=t)
    wb.save(AUD / "evidence-tracker.xlsx")


def build_pre_qual_checklist():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Pre-qualification")
    style_header_row(ws, ["#", "Requirement", "Owner", "Evidence", "Status", "Notes"])
    add_status_dropdown(ws, "E")
    rows = [
        ("PQ.1", "Microsoft AI Cloud Partner Program active"),
        ("PQ.2", "Solutions Partner - Digital & App Innovation, AND/OR Data & AI"),
        ("PQ.3", "Eligible Azure consumption threshold met"),
        ("PQ.4", "Required certified individuals - see Module A"),
        ("PQ.5", "Customer references aligned to Agentic DevOps"),
        ("PQ.6", "Audit engagement scheduled with Microsoft-approved auditor"),
    ]
    for i, (n, t) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=n)
        ws.cell(row=i, column=2, value=t)
    wb.save(AUD / "pre-qual-checklist.xlsx")


def main():
    build_qualification_questionnaire()
    build_definition_of_done()
    build_hld()
    build_lld()
    build_runbook()
    build_kt_plan()
    build_hypercare_plan()
    build_offering_one_pager()
    build_discovery_workshop_deck()
    build_discovery_workshop_workbook()
    build_waf_assessment()
    build_assessment_platform_inputs()
    build_evidence_tracker()
    build_pre_qual_checklist()
    files = sorted(ENG.glob("*")) + sorted(DEL.glob("*")) + sorted(AUD.glob("*"))
    print(f"Generated {len(files)} workfiles")
    for f in files:
        print(" -", f.relative_to(ROOT))


if __name__ == "__main__":
    main()
